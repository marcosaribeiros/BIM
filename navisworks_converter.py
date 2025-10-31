"""Convert Excel spreadsheets to Navisworks-compatible XML and XSD files."""

from __future__ import annotations

import argparse
import pathlib
from typing import List, Sequence

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import xml.etree.ElementTree as ET


XS_NS = "http://www.w3.org/2001/XMLSchema"
ET.register_namespace("xs", XS_NS)


def _indent(element: ET.Element, level: int = 0) -> None:
    """Indent an XML element tree in place for pretty output."""
    indentation = "\n" + "  " * level
    if len(element):
        if not element.text or not element.text.strip():
            element.text = indentation + "  "
        for child in element:
            _indent(child, level + 1)
            if not child.tail or not child.tail.strip():
                child.tail = indentation + "  "
        if not element[-1].tail or not element[-1].tail.strip():
            element[-1].tail = indentation
    elif level and (not element.tail or not element.tail.strip()):
        element.tail = indentation


def _coerce_to_text(value) -> str:
    """Return a string representation for a cell value."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        # Avoid scientific notation and preserve integer formatting
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return ("%f" % value).rstrip("0").rstrip(".")
    return str(value)


def _headers_from_sheet(sheet: Worksheet) -> List[str]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = []
    for index, value in enumerate(rows[0], start=1):
        header = _coerce_to_text(value) or f"Column{index}"
        headers.append(header)
    return headers


def _append_sheet_to_xml(parent: ET.Element, sheet: Worksheet) -> None:
    sheet_element = ET.SubElement(parent, "Sheet", name=sheet.title)
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return

    headers = _headers_from_sheet(sheet)
    for data_index, row in enumerate(rows[1:], start=1):
        row_element = ET.SubElement(sheet_element, "Row", index=str(data_index))
        for column_index, header in enumerate(headers):
            value = _coerce_to_text(row[column_index] if column_index < len(row) else None)
            field = ET.SubElement(row_element, "Field", name=header, order=str(column_index + 1))
            field.text = value


def workbook_to_xml(workbook_path: pathlib.Path) -> ET.ElementTree:
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    root = ET.Element("NavisworksData", workbook=workbook_path.stem)
    for sheet in wb.worksheets:
        _append_sheet_to_xml(root, sheet)
    return ET.ElementTree(root)


def build_xsd_tree() -> ET.ElementTree:
    schema = ET.Element(ET.QName(XS_NS, "schema"))
    schema.set("elementFormDefault", "qualified")

    navisworks = ET.SubElement(schema, ET.QName(XS_NS, "element"), name="NavisworksData")
    navisworks_type = ET.SubElement(navisworks, ET.QName(XS_NS, "complexType"))
    navisworks_sequence = ET.SubElement(navisworks_type, ET.QName(XS_NS, "sequence"))

    sheet_element = ET.SubElement(
        navisworks_sequence,
        ET.QName(XS_NS, "element"),
        name="Sheet",
        minOccurs="0",
        maxOccurs="unbounded",
    )
    sheet_type = ET.SubElement(sheet_element, ET.QName(XS_NS, "complexType"))
    sheet_sequence = ET.SubElement(sheet_type, ET.QName(XS_NS, "sequence"))

    row_element = ET.SubElement(
        sheet_sequence,
        ET.QName(XS_NS, "element"),
        name="Row",
        minOccurs="0",
        maxOccurs="unbounded",
    )
    row_type = ET.SubElement(row_element, ET.QName(XS_NS, "complexType"))
    row_sequence = ET.SubElement(row_type, ET.QName(XS_NS, "sequence"))

    field_element = ET.SubElement(
        row_sequence,
        ET.QName(XS_NS, "element"),
        name="Field",
        minOccurs="0",
        maxOccurs="unbounded",
    )
    field_type = ET.SubElement(field_element, ET.QName(XS_NS, "complexType"))
    simple_content = ET.SubElement(field_type, ET.QName(XS_NS, "simpleContent"))
    extension = ET.SubElement(simple_content, ET.QName(XS_NS, "extension"), base="xs:string")
    ET.SubElement(
        extension,
        ET.QName(XS_NS, "attribute"),
        name="name",
        type="xs:string",
        use="required",
    )
    ET.SubElement(
        extension,
        ET.QName(XS_NS, "attribute"),
        name="order",
        type="xs:integer",
        use="required",
    )

    ET.SubElement(
        row_type,
        ET.QName(XS_NS, "attribute"),
        name="index",
        type="xs:integer",
        use="required",
    )
    ET.SubElement(
        sheet_type,
        ET.QName(XS_NS, "attribute"),
        name="name",
        type="xs:string",
        use="required",
    )
    ET.SubElement(
        navisworks_type,
        ET.QName(XS_NS, "attribute"),
        name="workbook",
        type="xs:string",
        use="required",
    )

    return ET.ElementTree(schema)


def save_tree(tree: ET.ElementTree, path: pathlib.Path) -> None:
    root = tree.getroot()
    _indent(root)
    if not root.tail or not root.tail.strip():
        root.tail = "\n"
    tree.write(path, xml_declaration=True, encoding="utf-8", method="xml")


def convert_excel_to_navisworks(excel_path: pathlib.Path, output_dir: pathlib.Path) -> None:
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    output_dir.mkdir(parents=True, exist_ok=True)

    xml_tree = workbook_to_xml(excel_path)
    xml_path = output_dir / f"{excel_path.stem}.xml"
    save_tree(xml_tree, xml_path)

    xsd_tree = build_xsd_tree()
    xsd_path = output_dir / f"{excel_path.stem}.xsd"
    save_tree(xsd_tree, xsd_path)


def parse_args(args: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert an Excel spreadsheet (.xlsx) into XML and XSD files compatible with Autodesk Navisworks.",
    )
    parser.add_argument("excel", type=pathlib.Path, help="Path to the source Excel .xlsx file")
    parser.add_argument(
        "--output",
        type=pathlib.Path,
        default=pathlib.Path.cwd(),
        help="Directory where the XML and XSD files will be written (default: current directory)",
    )
    return parser.parse_args(args)


def main(argv: Sequence[str] | None = None) -> None:
    options = parse_args(argv)
    convert_excel_to_navisworks(options.excel, options.output)


if __name__ == "__main__":
    main()
